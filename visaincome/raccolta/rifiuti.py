"""Scarica e legge i file ufficiali sui visti rifiutati e scrive visaincome/dati/rifiuti.json.

Uso (dalla radice del repository):
    python3 visaincome/raccolta/rifiuti.py            # scarica (o riusa i file già in raccolta/_grezzi/) e scrive il JSON
    python3 visaincome/raccolta/rifiuti.py --nuovo    # riscarica sempre i file ufficiali

Fonti (vedi NOTE.md, sezione "Visa refusal rates"):
1. Commissione Europea, DG Migration and Home Affairs: "Visa statistics for consulates" (un XLSX per anno),
   elencati sulla pagina PAGINA_UE. Si prendono gli ultimi due anni pubblicati (anno più recente + confronto).
   Foglio "Data for consulates": una riga per Stato Schengen × Paese in cui si trova il consolato × consolato.
   Tasso ufficiale di "not issued" (rifiutati + esami interrotti, art. 8(2) del Codice dei visti):
   not issued / (rilasciati, compresi MEV + LTV + not issued). Lo script lo ricalcola e lo confronta con la colonna del file.
2. Home Office (Regno Unito), Immigration System Statistics: "Entry clearance visa applications and outcomes detailed
   datasets" (XLSX, foglio Data_Vis_D02: anno, trimestre, nazionalità, gruppo di visto, esito, decisioni), trovato
   dall'API dei contenuti di GOV.UK. Tasso di rifiuto = rifiutati / (concessi + rifiutati), come il "grant rate"
   ufficiale (che esclude ritirati e decaduti). Controllo incrociato con la tabella Vis_06 del file di sintesi (ODS).
3. Stati Uniti: travel.state.gov risponde 403 / non raggiungibile dal nostro server (15/10/2026): nessun dato.

Se un controllo fallisce, il JSON esistente NON viene toccato (esce con codice 1).
I file grezzi stanno in raccolta/_grezzi/ (esclusi da git, si riscaricano con questo script).
"""
import hashlib
import html
import json
import re
import sys
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from datetime import date
from pathlib import Path

import openpyxl

RADICE = Path(__file__).resolve().parent.parent
GREZZI = Path(__file__).resolve().parent / "_grezzi"
USCITA = RADICE / "dati" / "rifiuti.json"
UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0 Safari/537.36"}

PAGINA_UE = "https://home-affairs.ec.europa.eu/policies/schengen/visa-policy/short-stay-visas-issued-schengen-countries_en"
METADATI_UE = ("https://home-affairs.ec.europa.eu/document/download/b4626219-fa37-4650-8216-0c2bda466c74_en"
               "?filename=visa_statistics_-_quality_and_metadata_reporting.pdf")
PAGINA_UK = "https://www.gov.uk/government/statistical-data-sets/immigration-system-statistics-data-tables"
API_UK = "https://www.gov.uk/api/content/government/statistical-data-sets/immigration-system-statistics-data-tables"
TITOLO_UK = "Entry clearance visa applications and outcomes detailed datasets"
TITOLO_UK_SINTESI = "Entry clearance visas summary tables"
GRUPPI_UK = ("Visitor", "Study", "Work", "Family", "Other")
ESITI_UK = ("Issued", "Refused", "Withdrawn", "Lapsed")
TRIMESTRI_UK = 8  # ultimi 8 trimestri: anno più recente (4) + anno precedente (4)
# Tipo di visto "Family: Partner" (colonna "Visa type" del foglio D02): coniugi e partner, dentro il gruppo Family.
# Si salva come voce in più accanto ai 5 gruppi (non entra nei totali "All", che sommano solo GRUPPI_UK).
TIPO_PARTNER = "Family: Partner"


def scarica(url, dest=None, nuovo=False):
    """Scarica url (o riusa dest se c'è già). Restituisce (bytes, intestazioni)."""
    if dest is not None and dest.exists() and not nuovo:
        meta = json.loads(dest.with_suffix(dest.suffix + ".json").read_text()) if dest.with_suffix(dest.suffix + ".json").exists() else {}
        return dest.read_bytes(), meta
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=180) as r:
        dati = r.read()
        meta = {"http": r.status, "content_disposition": r.headers.get("Content-Disposition", ""),
                "last_modified": r.headers.get("Last-Modified", ""), "scaricato_il": date.today().isoformat()}
    if dest is not None:
        dest.parent.mkdir(parents=True, exist_ok=True)
        dest.write_bytes(dati)
        dest.with_suffix(dest.suffix + ".json").write_text(json.dumps(meta))
    return dati, meta


def impronta(b):
    return hashlib.sha256(b).hexdigest()


# ---------------------------------------------------------------- Unione Europea (Schengen)
def file_ue(pagina_html):
    """Dalla tabella della pagina: {anno: url} dei file 'Visa statistics for consulates'."""
    out = {}
    for cella in re.findall(r'<td[^>]*data-ecl-table-header="Visa statistics for consulates"[^>]*>(.*?)</td>', pagina_html, re.S):
        m = re.search(r'<a href="([^"]+)"[^>]*>\s*(\d{4})\s*</a>', cella)
        if m:
            url = html.unescape(m.group(1))
            out[int(m.group(2))] = url if url.startswith("http") else "https://home-affairs.ec.europa.eu" + url
    return out


def leggi_ue(percorso, anno):
    wb = openpyxl.load_workbook(percorso, read_only=True, data_only=True)
    ws = wb["Data for consulates"]
    righe, errori = [], []
    intest = None
    for r in ws.iter_rows(values_only=True):
        if intest is None:
            intest = [str(x or "").strip() for x in r]
            attese = {0: "Schengen State", 1: "Country where consulate is located", 2: "Consulate", 8: "Uniform visas applied for",
                      9: "Total uniform visas issued (including MEV)", 12: "Total LTVs issued", 13: "Uniform visas not issued",
                      14: "Not issued rate for uniform visas"}
            for i, t in attese.items():
                if intest[i] != t:
                    errori.append(f"UE {anno}: colonna {i} = {intest[i]!r}, attesa {t!r}")
            continue
        if not r[0] or not r[1]:
            continue
        stato, paese, cons = str(r[0]).strip(), str(r[1]).strip(), str(r[2] or "").strip()
        dom, ril, ltv, non = (int(r[i] or 0) for i in (8, 9, 12, 13))
        dec = ril + ltv + non
        tasso = non / dec if dec else None
        uff = r[14]
        if uff is not None and tasso is not None and abs(uff - tasso) > 1e-9:
            errori.append(f"UE {anno}: tasso ricalcolato diverso dal file per {stato}/{paese}/{cons}: {tasso} vs {uff}")
        righe.append([stato, paese, cons, dom, ril, ltv, non])
    # Controllo: somme per Stato = foglio "Totals - Schengen State"; somme per Paese = "Totals - third country"
    per_stato = defaultdict(lambda: [0, 0, 0, 0])
    per_paese = defaultdict(lambda: [0, 0, 0, 0])
    for s, p, _c, dom, ril, ltv, non in righe:
        for d in (per_stato[s], per_paese[p]):
            d[0] += dom; d[1] += ril; d[2] += ltv; d[3] += non
    tot_stati = {}
    for r in wb["Totals - Schengen State"].iter_rows(values_only=True):
        if r[0] and isinstance(r[1], (int, float)) and str(r[0]).strip() in per_stato:
            tot_stati[str(r[0]).strip()] = [int(r[i] or 0) for i in (1, 2, 4, 5)]
    for s, v in per_stato.items():
        if tot_stati.get(s) != v:
            errori.append(f"UE {anno}: somme di {s} {v} diverse dal foglio dei totali {tot_stati.get(s)}")
    tot_paesi = {}
    for r in wb["Totals - third country"].iter_rows(values_only=True):
        if r[1] and isinstance(r[2], (int, float)) and r[0] is not None:
            tot_paesi[str(r[1]).strip()] = [int(r[i] or 0) for i in (2, 3, 5, 6)]
    for p, v in per_paese.items():
        if p in tot_paesi and tot_paesi[p] != v:
            errori.append(f"UE {anno}: somme di {p} {v} diverse dal foglio dei totali {tot_paesi[p]}")
    if len(tot_paesi) < 100 or len(per_stato) < 25:
        errori.append(f"UE {anno}: troppo pochi Paesi ({len(tot_paesi)}) o Stati ({len(per_stato)})")
    return righe, errori


# ---------------------------------------------------------------- Regno Unito
def allegati_uk():
    d = json.loads(scarica(API_UK)[0])
    att = d["details"].get("attachments", [])
    det = next(a for a in att if a.get("title", "").startswith(TITOLO_UK))
    sin = next(a for a in att if a.get("title", "").startswith(TITOLO_UK_SINTESI))
    return d.get("public_updated_at", ""), det, sin


def leggi_uk(percorso):
    wb = openpyxl.load_workbook(percorso, read_only=True)
    cover = [str(r[0]) for r in wb["Cover_Sheet"].iter_rows(values_only=True) if r and r[0]]
    pubblicato = next((c.split(":", 1)[1].strip() for c in cover if c.startswith("Published")), "")
    prossimo = next((c.split(":", 1)[1].strip() for c in cover if c.startswith("Next update")), "")
    periodo = next((c for c in cover if c.startswith("Year ending")), "")
    ws = wb["Data_Vis_D02"]
    intest = None
    agg = defaultdict(lambda: [0, 0, 0, 0])  # (trimestre, nazionalità, gruppo) -> esiti
    trimestri = set()
    for r in ws.iter_rows(values_only=True):
        if intest is None:
            if r and r[0] == "Year" and r[1] == "Quarter":
                intest = list(r)
                if intest[2] != "Nationality" or intest[4] != "Visa type group" or intest[8] != "Case outcome" or intest[9] != "Decisions":
                    raise SystemExit(f"UK: intestazioni inattese {intest}")
            continue
        if not r[1]:
            continue
        trimestri.add(r[1])
        if r[4] not in GRUPPI_UK or r[8] not in ESITI_UK:
            raise SystemExit(f"UK: gruppo o esito inattesi {r}")
        agg[(r[1], r[2], r[4])][ESITI_UK.index(r[8])] += int(r[9] or 0)
        if r[5] == TIPO_PARTNER:  # sotto-voce del gruppo Family (dal 28/10/2026): visti per il partner (coniuge)
            if r[4] != "Family":
                raise SystemExit(f"UK: {TIPO_PARTNER} fuori dal gruppo Family {r}")
            agg[(r[1], r[2], TIPO_PARTNER)][ESITI_UK.index(r[8])] += int(r[9] or 0)
    ultimi = sorted(trimestri)[-TRIMESTRI_UK:]
    for (q, n, g), v in agg.items():  # il partner non può superare il gruppo Family
        if g == TIPO_PARTNER and any(x > y for x, y in zip(v, agg[(q, n, "Family")])):
            raise SystemExit(f"UK: {n} {q}: partner {v} oltre il gruppo Family {agg[(q, n, 'Family')]}")
    naz = defaultdict(dict)
    for (q, n, g), v in agg.items():
        if q in ultimi:
            naz[n].setdefault(g, {})[q] = v
    return {"pubblicato": pubblicato, "prossimo": prossimo, "periodo": periodo, "trimestri": ultimi, "naz": naz}


def leggi_ods(dati, foglio):
    """Legge un foglio di un file ODS (solo libreria standard) come lista di righe di testo."""
    ns = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0", "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
          "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
    t = lambda k: "{%s}%s" % (ns[k.split(":")[0]], k.split(":")[1])  # noqa: E731
    import io
    radice = ET.fromstring(zipfile.ZipFile(io.BytesIO(dati)).read("content.xml"))
    for tab in radice.iter(t("table:table")):
        if tab.get(t("table:name")) != foglio:
            continue
        righe = []
        for r in tab.iter(t("table:table-row")):
            riga = []
            for c in r:
                if c.tag not in (t("table:table-cell"), t("table:covered-table-cell")):
                    continue
                rip = min(int(c.get(t("table:number-columns-repeated"), "1")), 30)
                v = c.get(t("office:value"))
                if v is None:
                    v = " ".join("".join(p.itertext()) for p in c.iter(t("text:p"))) or None
                riga += [v] * rip
            righe.append([x for x in riga if x is not None])
        return righe
    return []


def controllo_vis06(uk, ods):
    """I totali dei visti turistici (ultimi 4 trimestri) devono coincidere con la tabella ufficiale Vis_06."""
    errori, n = [], 0
    ultimi4 = uk["trimestri"][-4:]
    for r in leggi_ods(ods, "Vis_06"):
        if len(r) >= 7 and r[0].isdigit():
            nome, conc, rif = r[1], int(float(r[4])), int(float(r[6]))
            v = uk["naz"].get(nome, {}).get("Visitor", {})
            mio = [sum(v.get(q, [0, 0, 0, 0])[i] for q in ultimi4) for i in (0, 1)]
            if mio != [conc, rif]:
                errori.append(f"UK Vis_06: {nome} concessi/rifiutati {mio} contro {conc}/{rif}")
            n += 1
    if n < 10:
        errori.append(f"UK Vis_06: lette solo {n} righe")
    return errori, n


# ---------------------------------------------------------------- principale
def main():
    nuovo = "--nuovo" in sys.argv
    errori = []
    oggi = date.today().isoformat()
    pag, _ = scarica(PAGINA_UE)
    pag = pag.decode("utf-8")
    anni = file_ue(pag)
    if len(anni) < 2:
        raise SystemExit("UE: file per anno non trovati sulla pagina")
    agg_pag = re.search(r'"dateModified":\s*"(\d{4}-\d{2}-\d{2})', pag)
    ue = {}
    fonti = {}
    for anno in sorted(anni)[-2:]:
        b, meta = scarica(anni[anno], GREZZI / f"schengen-consolati-{anno}.xlsx", nuovo)
        righe, e = leggi_ue(GREZZI / f"schengen-consolati-{anno}.xlsx", anno)
        errori += e
        ue[str(anno)] = righe
        nome = re.search(r'filename="?([^";]+)', meta.get("content_disposition", "") or "")
        fonti[f"schengen_{anno}"] = {
            "titolo": f"Visa statistics for consulates {anno}", "editore": "European Commission, DG Migration and Home Affairs",
            "url": anni[anno], "pagina": PAGINA_UE, "formato": "XLSX",
            "file": nome.group(1) if nome else anni[anno].rsplit("filename=", 1)[-1], "byte": len(b), "sha256": impronta(b),
            "scaricato_il": meta.get("scaricato_il", oggi), "pagina_aggiornata": agg_pag.group(1) if agg_pag else "", "righe": len(righe)}
    agg_uk, det, sin = allegati_uk()
    b_uk, meta_uk = scarica(det["url"], GREZZI / "uk-entry-clearance-outcomes.xlsx", nuovo)
    uk = leggi_uk(GREZZI / "uk-entry-clearance-outcomes.xlsx")
    ods, _ = scarica(sin["url"], GREZZI / "uk-visas-summary.ods", nuovo)
    e, n06 = controllo_vis06(uk, ods)
    errori += e
    fonti["uk"] = {"titolo": det["title"], "editore": "Home Office, Immigration System Statistics", "url": det["url"], "pagina": PAGINA_UK,
                   "formato": "XLSX", "byte": len(b_uk), "sha256": impronta(b_uk), "scaricato_il": meta_uk.get("scaricato_il", oggi),
                   "pubblicato": uk["pubblicato"], "prossimo": uk["prossimo"], "periodo": uk["periodo"], "pagina_aggiornata": agg_uk[:10]}
    fonti["uk_sintesi"] = {"titolo": sin["title"], "editore": "Home Office, Immigration System Statistics", "url": sin["url"], "pagina": PAGINA_UK,
                           "formato": "ODS", "righe_vis06_controllate": n06}
    if errori:
        print("ERRORI: il file dei dati NON è stato aggiornato")
        for x in errori[:40]:
            print(" -", x)
        sys.exit(1)
    uscita = {"_nota": "Generato da raccolta/rifiuti.py: non modificare a mano. Righe Schengen: [Stato, Paese del consolato, consolato, "
                       "domande, visti uniformi rilasciati (con MEV), LTV rilasciati, non rilasciati]. UK: nazionalità -> gruppo -> trimestre -> "
                       "[concessi, rifiutati, ritirati, decaduti]; la voce 'Family: Partner' è una parte del gruppo Family (tipo di visto del partner).",
              "generato_il": oggi, "fonti": fonti, "schengen": ue,
              "uk": {"trimestri": uk["trimestri"], "naz": uk["naz"]}}
    USCITA.write_text(json.dumps(uscita, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Scritto {USCITA}: Schengen {', '.join(ue)} ({', '.join(str(len(v)) for v in ue.values())} righe); "
          f"UK {uk['trimestri'][0]}–{uk['trimestri'][-1]}, {len(uk['naz'])} nazionalità, pubblicato {uk['pubblicato']}; "
          f"Vis_06: {n06} righe coincidono.")


if __name__ == "__main__":
    main()
